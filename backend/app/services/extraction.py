from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class ExtractedPage:
    text: str
    page_number: int | None = None


class DocumentExtractor:
    def extract(self, path: str, file_type: str) -> list[ExtractedPage]:
        suffix = file_type.lower().lstrip('.')
        if suffix == 'txt':
            return [ExtractedPage(Path(path).read_text(encoding='utf-8'))]
        if suffix == 'pdf':
            import fitz
            with fitz.open(path) as pdf:
                return [ExtractedPage(page.get_text(), index + 1) for index, page in enumerate(pdf)]
        if suffix == 'docx':
            from docx import Document as WordDocument
            doc = WordDocument(path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            paragraphs.extend(' | '.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
            return [ExtractedPage('\n'.join(paragraphs))]
        raise ValueError(f'Unsupported document format: {suffix}')


class RecursiveChunker:
    def __init__(self, chunk_size: int = 1200, overlap: int = 180) -> None:
        self.chunk_size = chunk_size
        self.overlap = overlap

    def split(self, text: str) -> list[str]:
        text = text.strip()
        if not text:
            return []
        separators = ['\n\n', '\n', '. ', ' ', '']
        parts = [text]
        for sep in separators:
            if all(len(part) <= self.chunk_size for part in parts):
                break
            if not sep:
                parts = [piece for part in parts for piece in (part[i:i + self.chunk_size] for i in range(0, len(part), self.chunk_size))]
            else:
                parts = [piece for part in parts for piece in part.split(sep) if piece]
        chunks: list[str] = []
        current = ''
        for part in parts:
            separator = ' ' if current else ''
            if current and len(current) + len(separator) + len(part) > self.chunk_size:
                chunks.append(current)
                current = current[-self.overlap:] + (' ' if self.overlap else '') + part
            else:
                current += separator + part
        if current:
            chunks.append(current)
        return chunks
